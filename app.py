import io
import re
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
import streamlit as st
from docx import Document
from fpdf import FPDF
from pypdf import PdfReader
from scipy.stats import entropy
from sklearn.metrics.pairwise import cosine_similarity

st.set_page_config(
    page_title="SmartHire | Resume Matcher & Fit Predictor",
    page_icon="💼",
    layout="wide",
)

MODELS_DIR = Path("models")
DATA_DIR = Path("data/processed")

GLOBAL_TECH_SKILLS = [
    "python", "r", "sql", "machine learning", "deep learning", "nlp",
    "data analysis", "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch",
    "tableau", "power bi", "matplotlib", "seaborn", "statistics", "aws",
    "docker", "kubernetes", "git", "spark", "hadoop", "excel", "java", "spring boot"
]

@st.cache_resource
def load_artifacts():
    classifier = joblib.load(MODELS_DIR / "classifier.pkl")
    tfidf = joblib.load(MODELS_DIR / "tfidf_vectorizer.pkl")
    rec_tfidf = joblib.load(MODELS_DIR / "job_tfidf_vectorizer.pkl")
    fit_rf = joblib.load(MODELS_DIR / "fit_predictor.pkl")
    df_jobs = pd.read_csv(DATA_DIR / "jobs_clean.csv")
    return classifier, tfidf, rec_tfidf, fit_rf, df_jobs

try:
    classifier, tfidf, rec_tfidf, fit_rf, df_jobs = load_artifacts()
except Exception as e:
    st.error(f"Error loading model artifacts: {e}")
    st.stop()

def extract_text_from_file(uploaded_file) -> str:
    file_bytes = io.BytesIO(uploaded_file.read())
    filename = uploaded_file.name.lower()
    
    if filename.endswith(".pdf"):
        reader = PdfReader(file_bytes)
        pages_text = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages_text).strip()
    elif filename.endswith(".docx"):
        doc = Document(file_bytes)
        paragraphs = [para.text for para in doc.paragraphs]
        return "\n".join(paragraphs).strip()
    elif filename.endswith(".txt"):
        return file_bytes.getvalue().decode("utf-8", errors="ignore").strip()
    return ""

def clean_text(text: str) -> str:
    text = re.sub(r"http\S+\s*", " ", text)
    text = re.sub(r"[#@]\S+", " ", text)
    text = re.sub(r"[^A-Za-z0-9\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip().lower()

def extract_skills(text: str, skill_pool: list) -> set:
    found = set()
    cleaned = text.lower()
    for skill in skill_pool:
        pattern = r"\b" + re.escape(skill) + r"\b"
        if re.search(pattern, cleaned):
            found.add(skill)
    return found

def calculate_profile_entropy(probabilities: np.ndarray) -> float:
    valid_probs = probabilities[probabilities > 0]
    return float(entropy(valid_probs, base=2))

def get_top_k_domains(probabilities: np.ndarray, classes: np.ndarray, k: int = 3) -> list:
    sorted_idx = probabilities.argsort()[::-1][:k]
    return [
        {"domain": str(classes[i]), "probability": round(float(probabilities[i]) * 100, 2)}
        for i in sorted_idx
    ]

def evaluate_shortlist_potential(domain_proba: float, readiness_score: float, top_similarity: float):
    prob = (0.45 * domain_proba) + (0.35 * readiness_score) + (0.20 * top_similarity)
    prob = float(np.clip(prob, 0.0, 99.9))

    if prob >= 75.0:
        status = "Strong Contender"
        action = "Fast-track to Technical Interview"
    elif prob >= 55.0:
        status = "High-Potential Pivot"
        action = "Schedule Phone Screening"
    elif prob >= 40.0:
        status = "Upskill Candidate"
        action = "Retain for Associate Pipeline"
    else:
        status = "Out of Scope"
        action = "Archive Profile"

    return round(prob, 2), status, action

def generate_pdf_report(domain_pred: str, domain_proba: float, readiness_score: float, 
                        shortlist_p: float, status: str, action: str,
                        fit_tier: str, specialization_idx: float, entropy_val: float,
                        identified_skills: set, missing_skills: set, 
                        df_rec: pd.DataFrame) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "SmartHire - Candidate Evaluation Report", ln=True, align="C")
    pdf.set_font("Helvetica", "I", 10)
    pdf.cell(0, 6, "Probabilistic Hiring Intelligence & Role Suitability Assessment", ln=True, align="C")
    pdf.ln(6)
    
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "1. Executive Decision & Alignment", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, f"- Shortlist Likelihood: {shortlist_p:.1f}% ({status})", ln=True)
    pdf.cell(0, 6, f"- Recommended Action: {action}", ln=True)
    pdf.cell(0, 6, f"- Predicted Specialization: {domain_pred} (Certainty: {domain_proba:.2f}%)", ln=True)
    pdf.cell(0, 6, f"- Specialization Index: {specialization_idx:.1f}% (Shannon Entropy: {entropy_val:.2f} bits)", ln=True)
    pdf.cell(0, 6, f"- Candidate Tier: {fit_tier} | Taxonomy Coverage: {readiness_score:.1f}%", ln=True)
    pdf.ln(4)
    
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "2. Competency Analysis", ln=True)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 6, "Identified Skills:", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, ", ".join(sorted([s.title() for s in identified_skills])) if identified_skills else "None")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 6, "Recommended Upskilling:", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 6, ", ".join(sorted([s.title() for s in missing_skills])) if missing_skills else "None")
    pdf.ln(4)
    
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "3. Top Matched Roles", ln=True)
    pdf.set_font("Helvetica", "", 9)
    for idx, row in df_rec.head(5).iterrows():
        title = str(row.get("title", "N/A"))
        company = str(row.get("company", "N/A"))
        match_s = row.get("Match Score (%)", 0.0)
        fit_p = row.get("Fit Likelihood (%)", 0.0)
        pdf.cell(0, 5, f"{idx+1}. {title} | {company} - Match: {match_s}% (Fit Likelihood: {fit_p}%)", ln=True)
    
    return bytes(pdf.output())

# Streamlit App Navigation
st.title("💼 SmartHire Executive Engine")
st.markdown("Automated resume domain classification, role readiness analysis, and probabilistic shortlisting.")

st.sidebar.header("Navigation & Settings")
app_mode = st.sidebar.radio("Processing Mode:", ["Single Resume Analysis", "Batch Screening (Multi-Upload to Excel)"])
top_k_jobs = st.sidebar.slider("Recommended Jobs Count:", min_value=3, max_value=15, value=5)

rec_job_col = "clean_job_text" if "clean_job_text" in df_jobs.columns else "text"
job_tfidf_vectors = rec_tfidf.transform(df_jobs[rec_job_col].fillna("").astype(str))

# --- MODE 1: SINGLE RESUME ANALYSIS ---
if app_mode == "Single Resume Analysis":
    input_method = st.radio("Input Format:", ["Upload Document", "Direct Text Input"], horizontal=True)
    resume_raw_text = ""

    if input_method == "Upload Document":
        uploaded_file = st.file_uploader("Upload candidate resume", type=["pdf", "docx", "txt"])
        if uploaded_file is not None:
            resume_raw_text = extract_text_from_file(uploaded_file)
            if resume_raw_text:
                st.success(f"Parsed {uploaded_file.name} ({len(resume_raw_text.split())} words)")
                with st.expander("Inspect Parsed Document Text"):
                    st.write(resume_raw_text[:1500] + ("..." if len(resume_raw_text) > 1500 else ""))
            else:
                st.error("No extractable plain text found inside document.")
    else:
        resume_raw_text = st.text_area("Paste Resume Content:", height=220, placeholder="Paste resume text here...")

    if st.button("Analyze Profile", type="primary"):
        if not resume_raw_text.strip():
            st.warning("Please submit a valid resume input before analyzing.")
            st.stop()

        cleaned_resume = clean_text(resume_raw_text)

        # 1. Classification & Probabilities
        res_vec = tfidf.transform([cleaned_resume])
        all_probs = classifier.predict_proba(res_vec)[0]
        classes = classifier.classes_
        domain_pred = classifier.predict(res_vec)[0]
        domain_proba = float(np.max(all_probs)) * 100

        ambiguity_entropy = calculate_profile_entropy(all_probs)
        max_entropy = float(np.log2(len(classes)))
        specialization_index = max(0.0, (1.0 - (ambiguity_entropy / max_entropy)) * 100)
        top_3_roles = get_top_k_domains(all_probs, classes, k=3)

        # 2. Competencies & Readiness
        identified_skills = extract_skills(cleaned_resume, GLOBAL_TECH_SKILLS)
        missing_skills = set(GLOBAL_TECH_SKILLS) - identified_skills
        readiness_score = (len(identified_skills) / len(GLOBAL_TECH_SKILLS)) * 100

        # 3. Matching
        res_rec_vec = rec_tfidf.transform([cleaned_resume])
        similarities = cosine_similarity(res_rec_vec, job_tfidf_vectors).flatten()
        top_indices = similarities.argsort()[::-1][:top_k_jobs]

        rec_job_rows = []
        res_word_count = len(cleaned_resume.split())

        for idx in top_indices:
            job_record = df_jobs.iloc[idx]
            job_text = str(job_record[rec_job_col]).lower()
            job_skills = extract_skills(job_text, GLOBAL_TECH_SKILLS)
            shared_skills = identified_skills.intersection(job_skills)
            match_ratio = len(shared_skills) / max(len(job_skills), 1)

            fit_features = np.array([[similarities[idx], len(shared_skills), match_ratio, res_word_count, len(job_text.split())]])
            try:
                fit_probability = fit_rf.predict_proba(fit_features)[0][1] * 100
            except Exception:
                fit_probability = similarities[idx] * 100

            rec_job_rows.append({
                "title": job_record.get("title", "Not Specified"),
                "company": job_record.get("company", "Not Specified"),
                "location": job_record.get("location", "Not Specified"),
                "Match Score (%)": round(similarities[idx] * 100, 2),
                "Fit Likelihood (%)": round(fit_probability, 2)
            })

        df_recommended = pd.DataFrame(rec_job_rows)
        top_job_score = float(df_recommended.iloc[0]["Match Score (%)"]) if not df_recommended.empty else 0.0

        # 4. Shortlist Engine
        shortlist_p, status, action = evaluate_shortlist_potential(domain_proba, readiness_score, top_job_score)

        fit_tier = (
            "Tier 1: High Market Alignment" if readiness_score >= 60 
            else "Tier 2: Solid Mid-Level Profile" if readiness_score >= 35 
            else "Tier 3: Early-Stage Competency"
        )

        st.subheader("📊 Candidate Evaluation")
        k1, k2, k3, k4 = st.columns(4)
        k1.metric("Shortlist Likelihood", f"{shortlist_p:.1f}%")
        k2.metric("Predicted Domain", str(domain_pred))
        k3.metric("Taxonomy Coverage", f"{readiness_score:.1f}%")
        k4.metric("Specialization Index", f"{specialization_index:.1f}%")

        st.info(f"**Verdict**: {status} — **Recommended Action**: {action}")

        # 5. Probabilistic Breakdown
        st.subheader("🎲 Uncertainty & Role Distribution Analysis")
        p_col1, p_col2 = st.columns(2)
        with p_col1:
            st.markdown("**Posterior Domain Probabilities (Top 3)**")
            for item in top_3_roles:
                st.write(f"**{item['domain']}** ({item['probability']}%)")
                st.progress(float(item["probability"]) / 100.0)

        with p_col2:
            st.markdown("**Profile Entropy & Dispersion**")
            st.metric("Shannon Ambiguity Entropy", f"{ambiguity_entropy:.2f} bits", help=f"Theoretical maximum: {max_entropy:.2f} bits")
            if ambiguity_entropy < 2.2:
                st.success("Targeted profile with concentrated niche domain signals.")
            elif ambiguity_entropy < 3.2:
                st.info("Balanced cross-functional profile with multiple viable paths.")
            else:
                st.warning("High dispersion profile: skills are spread evenly across distinct domains.")

        st.markdown(
            f"""
            > **Executive Assessment Briefing**
            > 
            > * **Shortlist Status**: **{status}** ({shortlist_p:.1f}% probability) — {action}.
            > * **Candidate Tier**: **{fit_tier}** ({readiness_score:.1f}% taxonomy coverage across {len(identified_skills)} monitored skills).
            > * **Classification**: Specialization concentrated in **{domain_pred}** ({domain_proba:.1f}% certainty), with runner-up alignment in **{top_3_roles[1]['domain']}** ({top_3_roles[1]['probability']}%).
            > * **Priority Vacancy**: **{df_recommended.iloc[0]['title']}** ({df_recommended.iloc[0]['Match Score (%)']}% content similarity).
            > * **Target Upskilling**: `{", ".join(list(missing_skills)[:4])}`.
            """
        )

        pdf_report_bytes = generate_pdf_report(
            domain_pred, domain_proba, readiness_score, shortlist_p, status, action,
            fit_tier, specialization_index, ambiguity_entropy, identified_skills, missing_skills, df_recommended
        )
        st.download_button("📥 Download Candidate Assessment Report (PDF)", pdf_report_bytes, f"SmartHire_Report_{domain_pred.replace(' ', '_')}.pdf", "application/pdf")
        st.markdown("---")

        tab_jobs, tab_skills = st.tabs(["🎯 Top Recommended Jobs", "🔍 Skill Gap Breakdown"])
        with tab_jobs:
            st.dataframe(df_recommended, use_container_width=True)
        with tab_skills:
            cid, cmiss = st.columns(2)
            with cid:
                st.success(f"**Identified Competencies ({len(identified_skills)})**")
                for s in sorted(identified_skills):
                    st.markdown(f"- {s.title()}")
            with cmiss:
                st.warning(f"**Recommended Skills to Acquire ({len(missing_skills)})**")
                for s in sorted(missing_skills):
                    st.markdown(f"- {s.title()}")

# --- MODE 2: BATCH SCREENING (MULTI-RESUME TO EXCEL) ---
else:
    st.subheader("📁 Batch Screening & Shortlist Matrix Ranking")
    st.markdown("Upload multiple candidate resumes (`.pdf`, `.docx`, `.txt`) to compute probabilistic conversion scores, rank candidates, and export an Excel decision matrix.")

    uploaded_files = st.file_uploader(
        "Upload Multiple Resumes:",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True
    )

    if uploaded_files:
        st.info(f"Loaded {len(uploaded_files)} resumes ready for processing.")

    if st.button("Run Batch Evaluation", type="primary"):
        if not uploaded_files:
            st.warning("Please upload at least one resume file to run batch analysis.")
            st.stop()

        batch_records = []
        progress_bar = st.progress(0)
        status_text = st.empty()

        for i, file in enumerate(uploaded_files):
            status_text.text(f"Evaluating ({i+1}/{len(uploaded_files)}): {file.name}")
            raw_text = extract_text_from_file(file)

            if not raw_text.strip():
                batch_records.append({
                    "Filename": file.name,
                    "Status": "Parsing Failed / Empty Text",
                    "Shortlist Probability (%)": 0.0,
                    "Candidate Status": "Out of Scope",
                    "Recommended Action": "Archive Profile",
                    "Primary Domain": "N/A",
                    "Domain Certainty (%)": 0.0,
                    "Runner-up Domain": "N/A",
                    "Runner-up Certainty (%)": 0.0,
                    "Shannon Entropy (bits)": 0.0,
                    "Specialization Index (%)": 0.0,
                    "Readiness Score (%)": 0.0,
                    "Candidate Tier": "N/A",
                    "Skills Count": 0,
                    "Identified Skills": "",
                    "Top Matched Job": "",
                    "Job Match Score (%)": 0.0
                })
                continue

            cleaned = clean_text(raw_text)

            # Inferences
            res_vec = tfidf.transform([cleaned])
            all_probs = classifier.predict_proba(res_vec)[0]
            classes = classifier.classes_
            domain = classifier.predict(res_vec)[0]
            confidence = float(np.max(all_probs)) * 100

            top_roles = get_top_k_domains(all_probs, classes, k=2)
            entropy_val = calculate_profile_entropy(all_probs)
            max_ent = float(np.log2(len(classes)))
            spec_idx = max(0.0, (1.0 - (entropy_val / max_ent)) * 100)

            matched_skills = extract_skills(cleaned, GLOBAL_TECH_SKILLS)
            readiness = (len(matched_skills) / len(GLOBAL_TECH_SKILLS)) * 100

            res_rec_vec = rec_tfidf.transform([cleaned])
            similarities = cosine_similarity(res_rec_vec, job_tfidf_vectors).flatten()
            top_idx = similarities.argmax()
            top_job = df_jobs.iloc[top_idx].get("title", "Not Specified")
            top_score = round(float(similarities[top_idx]) * 100, 2)

            tier = (
                "Tier 1 (High)" if readiness >= 60
                else "Tier 2 (Mid)" if readiness >= 35
                else "Tier 3 (Developing)"
            )

            shortlist_p, status, action = evaluate_shortlist_potential(confidence, readiness, top_score)

            batch_records.append({
                "Filename": file.name,
                "Status": "Success",
                "Shortlist Probability (%)": shortlist_p,
                "Candidate Status": status,
                "Recommended Action": action,
                "Primary Domain": str(domain),
                "Domain Certainty (%)": round(confidence, 2),
                "Runner-up Domain": top_roles[1]["domain"] if len(top_roles) > 1 else "N/A",
                "Runner-up Certainty (%)": top_roles[1]["probability"] if len(top_roles) > 1 else 0.0,
                "Shannon Entropy (bits)": round(entropy_val, 2),
                "Specialization Index (%)": round(spec_idx, 1),
                "Readiness Score (%)": round(readiness, 2),
                "Candidate Tier": tier,
                "Skills Count": len(matched_skills),
                "Identified Skills": ", ".join(sorted([s.title() for s in matched_skills])),
                "Top Matched Job": top_job,
                "Job Match Score (%)": top_score
            })

            progress_bar.progress((i + 1) / len(uploaded_files))

        status_text.text("Batch Processing Complete!")
        df_batch = pd.DataFrame(batch_records)
        df_batch = df_batch.sort_values(by="Shortlist Probability (%)", ascending=False).reset_index(drop=True)

        df_batch["Cohort Percentile"] = [
            f"Top {max(1, int(round((idx + 1) / len(df_batch) * 100)))}%"
            for idx in range(len(df_batch))
        ]

        st.success(f"Evaluated and ranked {len(df_batch)} candidates by shortlist probability.")
        st.dataframe(df_batch, use_container_width=True)

        excel_buffer = io.BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            df_batch.to_excel(writer, index=False, sheet_name="Shortlist_Rankings")

        st.download_button(
            label="📊 Download Batch Shortlist Analysis (.xlsx)",
            data=excel_buffer.getvalue(),
            file_name="SmartHire_Shortlist_Decision_Matrix.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )